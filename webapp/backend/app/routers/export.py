"""Export routes: PPTX cluster-detection, export-panel settings, and PPTX
generation."""
from __future__ import annotations
import base64
import io
import re

from fastapi import APIRouter, HTTPException, Response
from pptx import Presentation
from pptx.util import Inches, Pt

from diagen.clusters import find_clusters

from ..diagram_helpers import _get_session
from ..schemas import (
    ClusterDetectResponse, SessionIdRequest, ExportSettingsRequest, PptxGenerateRequest,
)

router = APIRouter()

_UNSAFE_FILENAME_CHARS = re.compile(r'[\\/:*?"<>|\r\n]+')


def _safe_filename(name: str, fallback: str = 'diagram') -> str:
    cleaned = _UNSAFE_FILENAME_CHARS.sub('_', (name or '').strip())
    return (cleaned or fallback)[:150]


@router.post("/api/export/pptx/detect", response_model=ClusterDetectResponse)
def export_pptx_detect(req: SessionIdRequest):
    """Pure analysis, no file I/O: connected-components + orphan
    suggestions for the active diagram, for the Export panel to render
    before the user picks anything — see diagen/clusters.py."""
    session = _get_session(req.session_id)
    result = find_clusters(session.diagram)
    return ClusterDetectResponse(**result)


@router.post("/api/export/settings")
def export_settings_save(req: ExportSettingsRequest):
    """Opaque settings blob owned entirely by the frontend's Export panel
    (slide assignments, sizes, positions, orphan resolutions) — persisted
    on the session so the next project save carries it along (see
    projects_save/projects_load in app/routers/projects.py)."""
    session = _get_session(req.session_id)
    session.export_settings = req.settings
    return {"ok": True}


@router.post("/api/export/pptx/generate")
def export_pptx_generate(req: PptxGenerateRequest):
    """Builds a .pptx in memory from PNGs the frontend already rasterized
    client-side (see ExportPanel.tsx — the live SVG is what's actually
    cropped/rasterized, so this never needs a server-side SVG renderer).
    `add_picture` embeds each image's bytes directly into the file, not a
    path, so it opens identically on any machine."""
    prs = Presentation()
    prs.slide_width = Inches(req.slide_width_in)
    prs.slide_height = Inches(req.slide_height_in)
    blank_layout = prs.slide_layouts[6]
    max_slide_number = max(s.slide_number for s in req.slides)
    for _ in range(max_slide_number):
        prs.slides.add_slide(blank_layout)

    for slide_req in req.slides:
        slide = prs.slides[slide_req.slide_number - 1]
        for placement in slide_req.placements:
            try:
                png_bytes = base64.b64decode(placement.image_png_base64, validate=True)
            except Exception as exc:
                raise HTTPException(400, f"Invalid image data: {exc}") from exc
            slide.shapes.add_picture(
                io.BytesIO(png_bytes),
                Inches(placement.x), Inches(placement.y),
                Inches(placement.w), Inches(placement.h),
            )
        # Phase 2 wires the UI to these; the endpoint already accepts them
        # so Phase 2 is a frontend-only change.
        for block, default_pt in ((slide_req.header, 28.0), (slide_req.content, 16.0)):
            if block is None:
                continue
            box = slide.shapes.add_textbox(Inches(block.x), Inches(block.y), Inches(block.w), Inches(block.h))
            tf = box.text_frame
            tf.word_wrap = True
            lines = block.text.split('\n') or ['']
            tf.text = lines[0]
            tf.paragraphs[0].font.size = Pt(block.size or default_pt)
            tf.paragraphs[0].font.name = block.font
            for line in lines[1:]:
                p = tf.add_paragraph()
                p.text = line
                p.font.size = Pt(block.size or default_pt)
                p.font.name = block.font

    buf = io.BytesIO()
    prs.save(buf)
    filename = _safe_filename(req.filename) + '.pptx'
    return Response(
        content=buf.getvalue(),
        media_type='application/vnd.openxmlformats-officedocument.presentationml.presentation',
        headers={'Content-Disposition': f'attachment; filename="{filename}"'},
    )
