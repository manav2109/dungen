"""Unit tests for diagen/pptx_shapes.py — the native-PowerPoint-shapes PPTX
generator (AutoShapes + real connectors, as opposed to app/main.py's
raster-picture generate path, covered separately in test_api_export.py)."""
from __future__ import annotations
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE, MSO_SHAPE_TYPE

from diagen.model import Node, Edge, Diagram
from diagen.spec import Spec
from diagen.defaults import DEFAULT_SHAPES
from diagen.pptx_shapes import _mso_shape_kind, render_native_placement


def _fixture():
    """One plain shape, one container with a child, one rows (ERD) node,
    and an edge between two plain shapes — enough to exercise every branch
    in render_native_placement."""
    spec = Spec.blank(use_defaults=True)
    d = Diagram(id='D1', scenario='', title='t')
    spec.diagrams['D1'] = d

    n1 = Node(id='n1', diagram='D1', parent='', type='process', label='Start',
              x=0, y=0, w=140, h=50)
    n2 = Node(id='n2', diagram='D1', parent='', type='decision', label='Check?',
              x=200, y=0, w=140, h=80)
    grp = Node(id='g1', diagram='D1', parent='', type='group', label='Group',
               x=0, y=150, w=300, h=200)
    child = Node(id='c1', diagram='D1', parent='g1', type='server', label='Server1',
                 x=20, y=190, w=120, h=60)
    ent = Node(id='ent1', diagram='D1', parent='', type='entity', label='User',
               x=400, y=0, w=180, h=30, meta='rows=id;name;email')
    for n in (n1, n2, grp, child, ent):
        d.nodes[n.id] = n
    d.edges = [
        Edge(id='e1', diagram='D1', source='n1', target='n2',
             relation='sequence_flow', label='go'),
    ]
    spec._link()
    return spec, d


def test_shape_kind_mapping_never_raises_for_any_default_entity_type():
    for shape_def in DEFAULT_SHAPES.values():
        kind = _mso_shape_kind(shape_def.shape)
        assert kind is not None


def test_shape_kind_mapping_covers_the_distinctive_kinds():
    assert _mso_shape_kind('diamond') == MSO_SHAPE.DIAMOND
    assert _mso_shape_kind('cylinder') == MSO_SHAPE.CAN
    assert _mso_shape_kind('parallelogram') == MSO_SHAPE.PARALLELOGRAM
    assert _mso_shape_kind('cloud') == MSO_SHAPE.CLOUD
    assert _mso_shape_kind('3d box') == MSO_SHAPE.CUBE
    assert _mso_shape_kind('chevron') == MSO_SHAPE.CHEVRON
    assert _mso_shape_kind('circle') == MSO_SHAPE.OVAL
    assert _mso_shape_kind('rounded rectangle') == MSO_SHAPE.ROUNDED_RECTANGLE
    assert _mso_shape_kind('stadium') == MSO_SHAPE.ROUNDED_RECTANGLE
    # Deliberately unmapped, same as render_drawio.py's SHAPE_STYLE —
    # falls back to a plain rectangle, not a KeyError.
    assert _mso_shape_kind('rectangle') == MSO_SHAPE.RECTANGLE
    assert _mso_shape_kind('title band') == MSO_SHAPE.RECTANGLE


def test_render_native_placement_produces_expected_shape_types(tmp_path):
    spec, d = _fixture()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_native_placement(slide, d, spec, ['n1', 'n2', 'g1', 'ent1'], (0.5, 0.5, 12.0, 6.0))

    out = tmp_path / 'native.pptx'
    prs.save(str(out))
    prs2 = Presentation(str(out))
    shapes = list(prs2.slides[0].shapes)

    auto_shapes = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE]
    # n1 (process), n2 (decision), g1 (container), c1 (server, pulled in as
    # g1's descendant even though not named explicitly), ent1's outer box.
    assert len(auto_shapes) == 5
    labels = {s.text for s in auto_shapes}
    assert {'Start', 'Check?', 'Group', 'Server1'} <= labels

    connectors = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.LINE]
    assert len(connectors) >= 1  # the n1->n2 edge, plus ent1's row divider

    text_boxes = [s for s in shapes if s.shape_type == MSO_SHAPE_TYPE.TEXT_BOX]
    row_texts = {s.text for s in text_boxes}
    assert {'id', 'name', 'email', 'go'} <= row_texts


def test_render_native_placement_excludes_a_container_child_when_not_selected():
    spec, d = _fixture()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Only n1 and n2 selected — the container, its child, and the rows node
    # must not appear at all.
    render_native_placement(slide, d, spec, ['n1', 'n2'], (0.5, 0.5, 12.0, 6.0))
    labels = {s.text for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.AUTO_SHAPE}
    assert labels == {'Start', 'Check?'}


def test_render_native_placement_handles_empty_selection_without_error():
    spec, d = _fixture()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_native_placement(slide, d, spec, [], (0.5, 0.5, 12.0, 6.0))
    assert len(list(slide.shapes)) == 0


def test_render_native_placement_sets_explicit_text_color_on_every_run(tmp_path):
    """Regression test: add_shape()'s default <p:style> resolves an
    unstyled run's font color to the theme's lt1 (white in the standard
    Office theme), so every label was rendering invisible-to-washed-out
    against a light fill until _apply_text started setting font.color.rgb
    explicitly. Round-trips through a real save/load, since python-pptx's
    in-memory Font.color can read back as None before the XML is
    serialized even when the attribute will be written correctly."""
    spec, d = _fixture()
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    render_native_placement(slide, d, spec, ['n1', 'n2', 'g1', 'ent1'], (0.5, 0.5, 12.0, 6.0))

    out = tmp_path / 'colors.pptx'
    prs.save(str(out))
    prs2 = Presentation(str(out))
    checked = 0
    for shape in prs2.slides[0].shapes:
        if not shape.has_text_frame:
            continue
        for paragraph in shape.text_frame.paragraphs:
            if not paragraph.runs:
                continue
            checked += 1
            color = paragraph.font.color
            assert color.type is not None, f"no explicit color on {paragraph.runs[0].text!r}"
            assert str(color.rgb) == '333333'
    assert checked > 0
